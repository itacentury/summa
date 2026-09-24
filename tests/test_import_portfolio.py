"""Unit tests for :mod:`scripts.import_portfolio_xlsx`."""

import io
import sqlite3
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, Final

import openpyxl
import pytest

from scripts.import_portfolio_xlsx import (
    DEFAULT_SHEET,
    ImportAbort,
    ImportSummary,
    PositionColumn,
    build_snapshot_rows,
    fill_band_map,
    find_position_columns,
    import_workbook,
    infer_kind,
    load_repaired_workbook,
    main,
    normalize_depot_name,
    parse_date_cell,
    parse_fx_overrides,
    parse_number,
    split_header,
    strip_style_extensions,
)
from summa.portfolio import value_eur

VALUE_COLUMN: Final[int] = 2
STYLES_PART: Final[str] = "xl/styles.xml"

# The layout of the private tracking workbook the importer was written for. The file is
# personal data and is deliberately not in the repo, so these constants are the only
# record of the shape scripts/import_portfolio_xlsx.py has to cope with.
REAL_POSITIONS: Final[tuple[tuple[int, str], ...]] = (
    (2, "MSCI World SRI EUR"),
    (5, "MSCI Europe ESG EUR"),
    (8, "Core Stoxx Europe 600 EUR"),
    (11, "FTSE All-World USD"),
    (14, "AMD"),
    (17, "Hensoldt"),
    (20, "Take-Two Interactive"),
    (23, "BGF Continental European Flex"),
    (26, "Deka-DividendenStrategie"),
    (29, "Deka-Industrie 4.0"),
)
REAL_LAST_COLUMN: Final[int] = 34
REAL_MERGES: Final[tuple[tuple[int, int], ...]] = ((1, 22), (23, 31), (32, 34))
REAL_BAND_LABELS: Final[dict[int, str]] = {
    1: "Trade Republic Depot",
    23: "Deka Depot",
    32: "Gesamt",
}


def _real_headers() -> dict[int, str]:
    """Return the workbook's row-2 headers, keyed by column index."""
    headers: dict[int, str] = {1: "Datum"}
    for column, header in REAL_POSITIONS:
        headers[column] = header
        headers[column + 1] = "Einzahlung"
        headers[column + 2] = "Delta"
    headers[32] = "Insgesamt"
    headers[33] = "Einzahlung gesamt"
    headers[34] = "Delta"
    return headers


def _real_bands() -> dict[int, str]:
    """Return the workbook's depot band per column."""
    return fill_band_map(REAL_BAND_LABELS, REAL_MERGES, REAL_LAST_COLUMN)


def _position(currency: str = "EUR") -> PositionColumn:
    """Return a minimal position sitting in the test value column."""
    return PositionColumn(
        column=VALUE_COLUMN,
        depot="Trade Republic",
        name="Example",
        currency=currency,
        kind="etf",
        sort_order=0,
    )


def _grid(
    values: list[float | None], deposits: list[float | None]
) -> tuple[list[str], list[dict[int, float | None]], list[dict[int, float | None]]]:
    """Build the dates and per-column mappings build_snapshot_rows expects."""
    dates: list[str] = [f"2025-11-{10 + index:02d}" for index in range(len(values))]
    return (
        dates,
        [{VALUE_COLUMN: value} for value in values],
        [{VALUE_COLUMN: deposit} for deposit in deposits],
    )


# --- Pure helpers -----------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (None, None),
        ("", None),
        ("   ", None),
        (True, None),
        (0, 0.0),
        (1131.16, 1131.16),
        (510.769999999999982, 510.77),
        (174.63999999999996, 174.64),
        (-25.5, -25.5),
        ("1.131,16 €", 1131.16),
        ("1 131,16 €", 1131.16),
        ("1131.16", 1131.16),
        ("abc", None),
        (datetime(2025, 11, 13), None),
    ],
)
def test_parse_number(raw: Any, expected: float | None) -> None:
    """parse_number rounds to cents and maps anything unusable to None."""
    assert parse_number(raw) == expected


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (datetime(2025, 11, 13), "2025-11-13"),
        (date(2026, 9, 14), "2026-09-14"),
        ("13.11.2025", "2025-11-13"),
        ("13.11.25", "2025-11-13"),
        ("  13.11.25  ", "2025-11-13"),
        ("Datum", None),
        ("", None),
        (45974, None),
        (None, None),
    ],
)
def test_parse_date_cell(raw: Any, expected: str | None) -> None:
    """parse_date_cell normalizes real dates and German date strings to ISO."""
    assert parse_date_cell(raw) == expected


@pytest.mark.parametrize(
    ("header", "expected_name", "expected_currency"),
    [
        ("MSCI World SRI EUR", "MSCI World SRI", "EUR"),
        ("Core Stoxx Europe 600 EUR", "Core Stoxx Europe 600", "EUR"),
        ("FTSE All-World USD", "FTSE All-World", "USD"),
        ("  Nvidia   usd ", "Nvidia", "USD"),
        ("AMD", "AMD", "EUR"),
        ("USD", "USD", "EUR"),
        ("Hensoldt", "Hensoldt", "EUR"),
        ("Take-Two Interactive", "Take-Two Interactive", "EUR"),
        ("Deka-Industrie 4.0", "Deka-Industrie 4.0", "EUR"),
        ("BGF Continental European Flex", "BGF Continental European Flex", "EUR"),
    ],
)
def test_split_header(header: str, expected_name: str, expected_currency: str) -> None:
    """split_header strips only a trailing currency code, never the whole name."""
    assert split_header(header) == (expected_name, expected_currency)


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("MSCI World SRI", "etf"),
        ("Core Stoxx Europe 600", "etf"),
        ("FTSE All-World", "etf"),
        ("iShares Core MSCI World UCITS", "etf"),
        ("Deka MSCI World", "etf"),
        ("Deka-Industrie 4.0", "fund"),
        ("BGF Continental European Flex", "fund"),
        ("AMD", "stock"),
        ("Take-Two Interactive", "stock"),
    ],
)
def test_infer_kind(name: str, expected: str) -> None:
    """infer_kind reads the markers in a name, ETFs winning over funds."""
    assert infer_kind(name) == expected


@pytest.mark.parametrize(
    ("band", "expected"),
    [
        ("Trade Republic Depot", "Trade Republic"),
        ("Deka Depot", "Deka"),
        ("  Deka   Depot ", "Deka"),
        ("trade republic depot", "trade republic"),
        ("Gesamt", "Gesamt"),
        ("Depot", "Depot"),
        ("", ""),
    ],
)
def test_normalize_depot_name(band: str, expected: str) -> None:
    """normalize_depot_name drops a trailing 'Depot' unless it is the whole label."""
    assert normalize_depot_name(band) == expected


def test_fill_band_map_spreads_merged_labels() -> None:
    """A merged band covers every column in its range, including column A."""
    bands: dict[int, str] = _real_bands()
    assert bands[1] == "Trade Republic Depot"
    assert bands[22] == "Trade Republic Depot"
    assert bands[23] == "Deka Depot"
    assert bands[31] == "Deka Depot"
    assert bands[32] == "Gesamt"


def test_fill_band_map_forward_fills_unmerged_labels() -> None:
    """Without merges a label carries to the right until the next one."""
    bands: dict[int, str] = fill_band_map(
        {2: "Trade Republic Depot", 5: "Deka Depot"}, [], 6
    )
    assert bands[1] == ""
    assert bands[4] == "Trade Republic Depot"
    assert bands[6] == "Deka Depot"


def test_find_position_columns_reads_the_real_layout() -> None:
    """Every position is found, and Datum, Delta and the total band are not."""
    columns: list[PositionColumn] = find_position_columns(
        _real_headers(), _real_bands()
    )

    assert [column.column for column in columns] == [
        column for column, _ in REAL_POSITIONS
    ]
    assert [column.depot for column in columns] == ["Trade Republic"] * 7 + ["Deka"] * 3
    assert [column.sort_order for column in columns] == [0, 1, 2, 3, 4, 5, 6, 0, 1, 2]
    assert ("FTSE All-World", "USD", "etf") == (
        columns[3].name,
        columns[3].currency,
        columns[3].kind,
    )
    assert ("AMD", "EUR", "stock") == (
        columns[4].name,
        columns[4].currency,
        columns[4].kind,
    )


@pytest.mark.parametrize(
    "headers",
    [
        {1: "Position", 2: "Einzahlung gesamt"},
        {1: "Position", 2: "Delta"},
        {1: "Position"},
        {1: "Einzahlung", 2: "Einzahlung"},
    ],
)
def test_find_position_columns_rejects_non_positions(headers: dict[int, str]) -> None:
    """Only a column headed beside an exact 'Einzahlung' starts a position."""
    bands: dict[int, str] = dict.fromkeys(headers, "Trade Republic Depot")
    assert find_position_columns(headers, bands) == []


def test_find_position_columns_rejects_a_position_without_a_band() -> None:
    """A position under no depot band aborts rather than inventing one."""
    with pytest.raises(ImportAbort, match="no depot band"):
        find_position_columns({1: "Orphan", 2: "Einzahlung"}, {1: "", 2: ""})


@pytest.mark.parametrize(
    ("pairs", "expected"),
    [
        ([], {}),
        (["USD=1.08"], {"USD": 1.08}),
        (["usd=1.08"], {"USD": 1.08}),
        ([" USD = 1.08 "], {"USD": 1.08}),
        (["USD=1.08", "CHF=0.94"], {"USD": 1.08, "CHF": 0.94}),
        (["EUR=1.0"], {"EUR": 1.0}),
    ],
)
def test_parse_fx_overrides(pairs: list[str], expected: dict[str, float]) -> None:
    """parse_fx_overrides normalizes the code and reads the rate."""
    assert parse_fx_overrides(pairs) == expected


@pytest.mark.parametrize(
    "pair",
    ["USD", "USD=", "=1.08", "USD=0", "USD=-1", "USD=abc", "US=1.08", "EUR=1.1"],
)
def test_parse_fx_overrides_rejects_bad_pairs(pair: str) -> None:
    """A malformed or impossible rate is a usage error, not a silent default."""
    with pytest.raises(ValueError):
        parse_fx_overrides([pair])


@pytest.mark.parametrize(
    ("styles", "expected"),
    [
        (b"<fill><patternFill /></fill>", b"<fill><patternFill /></fill>"),
        (
            b"<patternFill><extLst><ext uri='x'/></extLst></patternFill>",
            b"<patternFill></patternFill>",
        ),
        (b"<patternFill><extLst/></patternFill>", b"<patternFill></patternFill>"),
        (
            b"<a><extLst><x/></extLst></a><b><extLst><y/></extLst></b>",
            b"<a></a><b></b>",
        ),
    ],
)
def test_strip_style_extensions(styles: bytes, expected: bytes) -> None:
    """Every extLst element is removed and the rest of the stylesheet survives."""
    assert strip_style_extensions(styles) == expected


# --- The rule engine --------------------------------------------------------


def test_build_snapshot_rows_records_every_week_after_the_start() -> None:
    """A plain run writes one uncarried row per week."""
    dates, values, deposits = _grid([100.0, 110.0, 120.0], [50.0, None, None])
    rows, skipped = build_snapshot_rows([_position()], dates, values, deposits, {})

    assert skipped == 0
    assert [row.value for row in rows] == [100.0, 110.0, 120.0]
    assert [row.deposit for row in rows] == [50.0, 0.0, 0.0]
    assert not any(row.carried for row in rows)


@pytest.mark.parametrize(
    ("values", "deposits", "expected_values", "expected_skipped"),
    [
        ([None, None, 200.0], [None, None, 200.0], [200.0], 2),
        ([0, 0, 0, 512.4], [None, None, None, 500.0], [512.4], 3),
        ([None, None], [None, None], [], 2),
    ],
)
def test_build_snapshot_rows_skips_weeks_before_a_position_exists(
    values: list[float | None],
    deposits: list[float | None],
    expected_values: list[float],
    expected_skipped: int,
) -> None:
    """A position starts at its first real value, whether blank or a literal 0."""
    dates, value_rows, deposit_rows = _grid(values, deposits)
    rows, skipped = build_snapshot_rows(
        [_position()], dates, value_rows, deposit_rows, {}
    )

    assert [row.value for row in rows] == expected_values
    assert skipped == expected_skipped


@pytest.mark.parametrize("first_value", [0, None])
def test_build_snapshot_rows_starts_on_a_deposit_without_a_value(
    first_value: float | None,
) -> None:
    """Money moving starts the position even when the value cell is blank."""
    dates, values, deposits = _grid([first_value, 120.0], [500.0, None])
    rows, skipped = build_snapshot_rows([_position()], dates, values, deposits, {})

    assert skipped == 0
    assert (rows[0].value, rows[0].deposit, rows[0].carried) == (0.0, 500.0, False)


def test_build_snapshot_rows_does_not_carry_before_the_first_value() -> None:
    """A week with no value to copy yet is not carried -- there is nothing to copy."""
    dates, values, deposits = _grid(
        [None, None, 150.0, None], [500.0, None, None, None]
    )
    rows, _ = build_snapshot_rows([_position()], dates, values, deposits, {})

    assert [row.value for row in rows] == [0.0, 0.0, 150.0, 150.0]
    assert [row.carried for row in rows] == [False, False, False, True]


def test_build_snapshot_rows_carries_a_blank_week_forward() -> None:
    """A gap keeps the previous value and admits it with the carried flag."""
    dates, values, deposits = _grid([100.0, None, 130.0], [None, None, None])
    rows, _ = build_snapshot_rows([_position()], dates, values, deposits, {})

    assert [row.value for row in rows] == [100.0, 100.0, 130.0]
    assert [row.carried for row in rows] == [False, True, False]


def test_build_snapshot_rows_keeps_a_deposit_on_a_carried_week() -> None:
    """carried describes the value only — a deposit stays at its own date."""
    dates, values, deposits = _grid([100.0, None, 180.0], [None, 50.0, None])
    rows, _ = build_snapshot_rows([_position()], dates, values, deposits, {})

    assert (rows[1].value, rows[1].deposit, rows[1].carried) == (100.0, 50.0, True)


def test_build_snapshot_rows_carries_to_the_end_of_the_grid() -> None:
    """Every position has a row on every date at or after its start."""
    dates, values, deposits = _grid([100.0, None, None], [None, None, None])
    rows, _ = build_snapshot_rows([_position()], dates, values, deposits, {})

    assert len(rows) == len(dates)
    assert [row.carried for row in rows] == [False, True, True]


def test_build_snapshot_rows_keeps_a_zero_after_the_start() -> None:
    """A position wiped out after it started records the zero, it does not skip it."""
    dates, values, deposits = _grid([100.0, 0], [None, None])
    rows, skipped = build_snapshot_rows([_position()], dates, values, deposits, {})

    assert skipped == 0
    assert (rows[1].value, rows[1].carried) == (0.0, False)


def test_build_snapshot_rows_applies_the_rate_of_the_position_currency() -> None:
    """Only a position quoted in the overridden currency gets its rate."""
    dates, values, deposits = _grid([100.0], [None])
    rows, _ = build_snapshot_rows(
        [_position("USD")], dates, values, deposits, {"USD": 1.08}
    )
    assert rows[0].fx_rate == 1.08

    rows, _ = build_snapshot_rows(
        [_position("EUR")], dates, values, deposits, {"USD": 1.08}
    )
    assert rows[0].fx_rate == 1.0


# --- End to end -------------------------------------------------------------

WORKBOOK_HEADERS: Final[tuple[str, ...]] = (
    "Datum",
    "MSCI World SRI EUR",
    "Einzahlung",
    "Delta",
    "FTSE All-World USD",
    "Einzahlung",
    "Delta",
    "Deka-Industrie 4.0",
    "Einzahlung",
    "Delta",
    "Insgesamt",
    "Einzahlung gesamt",
    "Delta",
)
WORKBOOK_ROWS: Final[tuple[tuple[Any, ...], ...]] = (
    (
        datetime(2025, 11, 13),
        0,
        None,
        0,
        None,
        None,
        None,
        50.0,
        50.0,
        0,
        50.0,
        50.0,
        0,
    ),
    (
        datetime(2025, 11, 20),
        100.0,
        100.0,
        0,
        None,
        None,
        None,
        51.0,
        None,
        1.0,
        151.0,
        100.0,
        0,
    ),
    (
        datetime(2025, 11, 27),
        None,
        50.0,
        0,
        200.0,
        200.0,
        0,
        52.0,
        None,
        1.0,
        352.0,
        250.0,
        0,
    ),
    (
        datetime(2025, 12, 4),
        130.0,
        None,
        30.0,
        210.0,
        None,
        10.0,
        53.0,
        None,
        1.0,
        393.0,
        0,
        41.0,
    ),
)


def _write_workbook(path: Path, headers: tuple[str, ...] = WORKBOOK_HEADERS) -> Path:
    """Write a miniature depot workbook mirroring the real layout."""
    workbook: Any = openpyxl.Workbook()
    worksheet: Any = workbook.active
    worksheet.title = "Übersicht"
    worksheet["A1"] = "Trade Republic Depot"
    worksheet["H1"] = "Deka Depot"
    worksheet["K1"] = "Gesamt"
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    worksheet.merge_cells(start_row=1, start_column=8, end_row=1, end_column=10)
    worksheet.merge_cells(start_row=1, start_column=11, end_row=1, end_column=13)
    worksheet.append(list(headers))
    for row in WORKBOOK_ROWS:
        worksheet.append(list(row))
    workbook.save(path)
    return path


def _connect(path: Path) -> sqlite3.Connection:
    """Open an imported database for inspection."""
    conn: sqlite3.Connection = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    return conn


def test_main_imports_a_workbook(tmp_path: Path) -> None:
    """A full run creates the depots, the positions and their snapshots."""
    workbook: Path = _write_workbook(tmp_path / "book.xlsx")
    database: Path = tmp_path / "test.db"

    assert main([str(workbook), "--db", str(database)]) == 0

    with _connect(database) as conn:
        depots: list[Any] = list(
            conn.execute(
                "SELECT name, sort_order FROM portfolio_depots ORDER BY sort_order"
            )
        )
        assert [(row["name"], row["sort_order"]) for row in depots] == [
            ("Trade Republic", 0),
            ("Deka", 1),
        ]

        positions: list[Any] = list(
            conn.execute(
                "SELECT name, kind, currency FROM portfolio_positions ORDER BY depot_id, sort_order"
            )
        )
        assert [(row["name"], row["kind"], row["currency"]) for row in positions] == [
            ("MSCI World SRI", "etf", "EUR"),
            ("FTSE All-World", "etf", "USD"),
            ("Deka-Industrie 4.0", "fund", "EUR"),
        ]

        counts: list[Any] = list(
            conn.execute(
                "SELECT p.name, COUNT(s.id) AS total, MIN(s.date) AS first "
                "FROM portfolio_positions p JOIN portfolio_snapshots s ON s.position_id = p.id "
                "GROUP BY p.id ORDER BY p.id"
            )
        )
        assert [(row["name"], row["total"], row["first"]) for row in counts] == [
            ("MSCI World SRI", 3, "2025-11-20"),
            ("FTSE All-World", 2, "2025-11-27"),
            ("Deka-Industrie 4.0", 4, "2025-11-13"),
        ]


def test_main_carries_a_blank_week(tmp_path: Path) -> None:
    """The blank week keeps the previous value, its own deposit and the flag."""
    workbook: Path = _write_workbook(tmp_path / "book.xlsx")
    database: Path = tmp_path / "test.db"
    assert main([str(workbook), "--db", str(database)]) == 0

    with _connect(database) as conn:
        row: Any = conn.execute(
            "SELECT s.value, s.deposit, s.carried FROM portfolio_snapshots s "
            "JOIN portfolio_positions p ON p.id = s.position_id "
            "WHERE p.name = ? AND s.date = ?",
            ("MSCI World SRI", "2025-11-27"),
        ).fetchone()

    assert (row["value"], row["deposit"], row["carried"]) == (100.0, 50.0, 1)


def test_main_never_rewrites_a_recorded_week(tmp_path: Path) -> None:
    """A second run leaves a week corrected in the UI exactly as it was."""
    workbook: Path = _write_workbook(tmp_path / "book.xlsx")
    database: Path = tmp_path / "test.db"
    assert main([str(workbook), "--db", str(database)]) == 0

    with _connect(database) as conn:
        conn.execute(
            "UPDATE portfolio_snapshots SET value = 999.0 WHERE date = ?",
            ("2025-12-04",),
        )
        conn.commit()
        before: int = conn.execute(
            "SELECT COUNT(*) AS n FROM portfolio_snapshots"
        ).fetchone()["n"]

    assert main([str(workbook), "--db", str(database)]) == 0

    with _connect(database) as conn:
        after: int = conn.execute(
            "SELECT COUNT(*) AS n FROM portfolio_snapshots"
        ).fetchone()["n"]
        corrected: list[Any] = list(
            conn.execute(
                "SELECT value FROM portfolio_snapshots WHERE date = ?", ("2025-12-04",)
            )
        )

    assert after == before
    assert all(row["value"] == 999.0 for row in corrected)


def test_import_workbook_counts_carried_rows_it_wrote(tmp_path: Path) -> None:
    """A re-run writes nothing, so it must not go on claiming carried rows either."""
    workbook: Path = _write_workbook(tmp_path / "book.xlsx")
    database: Path = tmp_path / "test.db"

    first: ImportSummary = import_workbook(workbook, DEFAULT_SHEET, {}, database, False)
    again: ImportSummary = import_workbook(workbook, DEFAULT_SHEET, {}, database, False)

    assert 0 < first.rows_carried <= first.snapshots_written
    assert again.snapshots_written == 0
    assert again.rows_carried == 0


def test_main_applies_the_fx_rate(tmp_path: Path) -> None:
    """--fx lands on the position quoted in that currency and nowhere else."""
    workbook: Path = _write_workbook(tmp_path / "book.xlsx")
    database: Path = tmp_path / "test.db"

    assert main([str(workbook), "--db", str(database), "--fx", "USD=1.08"]) == 0

    with _connect(database) as conn:
        rates: list[Any] = list(
            conn.execute(
                "SELECT p.name, s.fx_rate, s.value FROM portfolio_snapshots s "
                "JOIN portfolio_positions p ON p.id = s.position_id "
                "WHERE s.date = ? ORDER BY p.id",
                ("2025-12-04",),
            )
        )

    by_name: dict[str, Any] = {row["name"]: row for row in rates}
    assert by_name["FTSE All-World"]["fx_rate"] == 1.08
    assert by_name["MSCI World SRI"]["fx_rate"] == 1.0
    assert value_eur(
        by_name["FTSE All-World"]["value"], by_name["FTSE All-World"]["fx_rate"]
    ) == pytest.approx(210.0 / 1.08)


def test_main_dry_run_does_not_create_the_database(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """A dry run against a missing database leaves the path missing."""
    workbook: Path = _write_workbook(tmp_path / "book.xlsx")
    database: Path = tmp_path / "test.db"

    assert main([str(workbook), "--db", str(database), "--dry-run"]) == 0

    assert "nothing written" in capsys.readouterr().out
    assert not database.exists()
    assert not list(tmp_path.glob("test.db*"))


def test_main_dry_run_leaves_an_existing_database_untouched(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """A second run as a dry run reports the weeks it would skip, byte for byte."""
    workbook: Path = _write_workbook(tmp_path / "book.xlsx")
    database: Path = tmp_path / "test.db"

    assert main([str(workbook), "--db", str(database)]) == 0
    capsys.readouterr()
    before: bytes = database.read_bytes()

    assert main([str(workbook), "--db", str(database), "--dry-run"]) == 0

    output: str = capsys.readouterr().out
    assert "snapshots      0 written" in output
    assert "(0 created)" in output
    assert database.read_bytes() == before


def test_main_rejects_a_missing_workbook(tmp_path: Path) -> None:
    """A path that is not a file is a usage error."""
    assert main([str(tmp_path / "nope.xlsx"), "--db", str(tmp_path / "test.db")]) == 2


def test_main_rejects_a_bad_fx_pair(tmp_path: Path) -> None:
    """A malformed --fx is rejected before the workbook is even opened."""
    workbook: Path = _write_workbook(tmp_path / "book.xlsx")
    assert main([str(workbook), "--db", str(tmp_path / "test.db"), "--fx", "USD"]) == 2


def test_main_reports_a_sheet_without_positions(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """A sheet with no 'Einzahlung' column fails loudly instead of importing nothing."""
    headers: tuple[str, ...] = tuple(
        "Beitrag" if header == "Einzahlung" else header for header in WORKBOOK_HEADERS
    )
    workbook: Path = _write_workbook(tmp_path / "book.xlsx", headers)

    assert main([str(workbook), "--db", str(tmp_path / "test.db")]) == 1
    assert "no positions found" in capsys.readouterr().err


def test_load_repaired_workbook_reads_a_stylesheet_openpyxl_rejects(
    tmp_path: Path,
) -> None:
    """The repair is load-bearing: the same file fails to open without it."""
    source: Path = _write_workbook(tmp_path / "book.xlsx")
    broken: Path = tmp_path / "broken.xlsx"

    buffer: io.BytesIO = io.BytesIO()
    with zipfile.ZipFile(source) as original:
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
            for item in original.infolist():
                data: bytes = original.read(item.filename)
                if item.filename == STYLES_PART:
                    data = data.replace(
                        b"<patternFill />",
                        b'<patternFill><extLst><ext uri="{X}"><smNativeData /></ext></extLst></patternFill>',
                        1,
                    )
                target.writestr(item, data)
    broken.write_bytes(buffer.getvalue())

    with pytest.raises(TypeError):
        openpyxl.load_workbook(broken)

    worksheet: Any = load_repaired_workbook(broken)["Übersicht"]
    assert worksheet["A1"].value == "Trade Republic Depot"
