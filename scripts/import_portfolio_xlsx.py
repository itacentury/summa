"""Import a depot workbook into Summa's portfolio tables.

The workbook is wide: row 1 carries merged depot bands, row 2 the headers, and
from row 3 every row is one week. Each position occupies a triple of columns --
its value, its ``Einzahlung`` and a ``Delta`` that this importer recomputes
rather than reads, because :func:`summa.portfolio.week_delta` derives it from
the stored snapshots anyway.

Re-running is safe: snapshots are written with ``INSERT OR IGNORE`` on
``(position_id, date)``, so a week already in the database is left exactly as it
is -- including one corrected by hand in the UI.

``--dry-run`` runs the whole import against an in-memory copy of the database, so
it reports real counts while never creating or modifying the file behind ``--db``.
"""

import argparse
import io
import math
import re
import sqlite3
import sys
import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Final

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from scripts.portfolio_db import connect, connect_mirror, default_database_path

BAND_ROW: Final[int] = 1
HEADER_ROW: Final[int] = 2
FIRST_DATA_ROW: Final[int] = 3
DATE_COLUMN: Final[int] = 1

DEPOSIT_HEADER: Final[str] = "Einzahlung"
BAND_SUFFIX: Final[str] = "Depot"
IGNORED_BANDS: Final[frozenset[str]] = frozenset({"Gesamt", "Insgesamt", "Total"})
DEFAULT_SHEET: Final[str] = "Übersicht"

DEFAULT_CURRENCY: Final[str] = "EUR"
CURRENCY_CODE_LENGTH: Final[int] = 3
DEFAULT_FX_RATE: Final[float] = 1.0
MONEY_DIGITS: Final[int] = 2
# Not the ISO table, just what a depot sheet plausibly quotes in. A header's last
# token is only read as a currency when it appears here, which is what keeps a
# position named "AMD" from being stripped down to nothing.
CURRENCY_CODES: Final[frozenset[str]] = frozenset(
    {"EUR", "USD", "GBP", "CHF", "JPY", "SEK", "NOK", "DKK", "CAD", "AUD"}
)

ETF_MARKERS: Final[tuple[str, ...]] = ("msci", "stoxx", "ftse", "ucits")
FUND_MARKERS: Final[tuple[str, ...]] = ("deka", "bgf")
KIND_ETF: Final[str] = "etf"
KIND_FUND: Final[str] = "fund"
KIND_STOCK: Final[str] = "stock"

STYLES_PART: Final[str] = "xl/styles.xml"
# PlanMaker writes style extensions openpyxl does not model, and a <extLst>
# inside a <patternFill> makes load_workbook raise TypeError outright.
STYLE_EXTENSIONS: Final[re.Pattern[bytes]] = re.compile(
    rb"<extLst\b[^>]*/>|<extLst\b[^>]*>.*?</extLst>", re.DOTALL
)

DATE_FORMATS: Final[tuple[str, ...]] = ("%d.%m.%Y", "%d.%m.%y")

EXIT_OK: Final[int] = 0
EXIT_IMPORT_ERROR: Final[int] = 1
EXIT_USAGE_ERROR: Final[int] = 2


class ImportAbort(Exception):
    """The workbook could not be understood and nothing was written."""


@dataclass(frozen=True)
class PositionColumn:
    """One position's value column, with the depot band it sits under."""

    column: int
    depot: str
    name: str
    currency: str
    kind: str
    sort_order: int


@dataclass(frozen=True)
class SnapshotRow:
    """One row destined for ``portfolio_snapshots``, still keyed by name."""

    depot: str
    position: str
    date: str
    value: float
    deposit: float
    fx_rate: float
    carried: bool


@dataclass(frozen=True)
class ImportSummary:
    """What a run did, for the printed report."""

    depots: int
    depots_created: int
    positions: int
    positions_created: int
    snapshots_written: int
    snapshots_existing: int
    rows_carried: int
    rows_skipped: int


# --- Pure helpers -----------------------------------------------------------


def strip_style_extensions(styles_xml: bytes) -> bytes:
    """Remove every ``extLst`` element from a stylesheet.

    :param styles_xml: the raw ``xl/styles.xml`` part.
    """
    return STYLE_EXTENSIONS.sub(b"", styles_xml)


def parse_number(raw: object) -> float | None:
    """Coerce a cell value to an amount in cents precision, or None when empty.

    A repaired workbook hands over real floats, so the German text branch only
    catches a sheet whose amounts were saved as strings. The rounding is what
    turns a stored ``510.769999999999982`` back into ``510.77``.
    """
    if raw is None or isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)):
        return round(float(raw), MONEY_DIGITS)
    if not isinstance(raw, str):
        return None

    text: str = raw.replace(" ", " ").replace("€", "").strip()
    if not text:
        return None
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    text = text.replace(" ", "")
    try:
        return round(float(text), MONEY_DIGITS)
    except ValueError:
        return None


def parse_date_cell(raw: object) -> str | None:
    """Coerce a cell value to an ISO date, or None when it is not one."""
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, date):
        return raw.isoformat()
    if not isinstance(raw, str):
        return None

    text: str = raw.strip()
    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format).date().isoformat()
        except ValueError:
            continue
    return None


def split_header(header: str) -> tuple[str, str]:
    """Split a row-2 header into a position name and its currency.

    The trailing token is only taken as a currency when the header has more than
    one token, so a position literally named ``AMD`` or ``USD`` keeps its name.
    """
    tokens: list[str] = header.split()
    if len(tokens) > 1 and tokens[-1].upper() in CURRENCY_CODES:
        return " ".join(tokens[:-1]), tokens[-1].upper()
    return " ".join(tokens), DEFAULT_CURRENCY


def normalize_depot_name(band: str) -> str:
    """Turn a row-1 band label into the depot name the UI shows."""
    tokens: list[str] = band.split()
    if len(tokens) > 1 and tokens[-1].casefold() == BAND_SUFFIX.casefold():
        tokens = tokens[:-1]
    return " ".join(tokens)


def infer_kind(name: str) -> str:
    """Guess a position's kind from its name.

    ETF markers win over fund markers, so a hypothetical "Deka MSCI World" reads
    as an ETF. The user can correct either one in Settings afterwards.
    """
    folded: str = name.casefold()
    if any(marker in folded for marker in ETF_MARKERS):
        return KIND_ETF
    if any(marker in folded for marker in FUND_MARKERS):
        return KIND_FUND
    return KIND_STOCK


def fill_band_map(
    row_one: Mapping[int, str],
    merges: Sequence[tuple[int, int]],
    last_column: int,
) -> dict[int, str]:
    """Map every column to the depot band above it.

    Merged ranges carry their label in the top-left cell only. Columns no merge
    covers inherit the nearest label to their left, so a sheet that repeats the
    label instead of merging maps identically.

    :param merges: ``(min_col, max_col)`` pairs of the row-1 merged ranges.
    """
    merged_labels: dict[int, str] = {}
    for min_col, max_col in merges:
        label: str = row_one.get(min_col, "").strip()
        if not label:
            continue
        for column in range(min_col, max_col + 1):
            merged_labels[column] = label

    bands: dict[int, str] = {}
    current: str = ""
    for column in range(1, last_column + 1):
        current = (
            merged_labels.get(column) or row_one.get(column, "").strip() or current
        )
        bands[column] = current
    return bands


def find_position_columns(
    headers: Mapping[int, str], bands: Mapping[int, str]
) -> list[PositionColumn]:
    """Find the value column of every position in the sheet.

    A column starts a position exactly when the column beside it is headed
    ``Einzahlung``. That single test also rejects the ``Datum`` column, every
    ``Delta`` column and the sheet's own total band, whose deposit column is
    headed ``Einzahlung gesamt`` -- which is why the comparison is equality and
    not a prefix match.

    :raises ImportAbort: a position sits under no depot band.
    """
    columns: list[PositionColumn] = []
    order_per_depot: dict[str, int] = {}
    for column in sorted(headers):
        header: str = headers[column].strip()
        if not header or header == DEPOSIT_HEADER:
            continue
        if headers.get(column + 1, "").strip() != DEPOSIT_HEADER:
            continue

        depot: str = normalize_depot_name(bands.get(column, ""))
        if depot in IGNORED_BANDS:
            continue
        if not depot:
            raise ImportAbort(
                f"column {get_column_letter(column)} ({header!r}) has no depot band in row {BAND_ROW}"
            )

        name, currency = split_header(header)
        sort_order: int = order_per_depot.get(depot, 0)
        order_per_depot[depot] = sort_order + 1
        columns.append(
            PositionColumn(
                column=column,
                depot=depot,
                name=name,
                currency=currency,
                kind=infer_kind(name),
                sort_order=sort_order,
            )
        )
    return columns


def _position_started(value: float | None, deposit: float | None) -> bool:
    """Decide whether a position is already held in this row.

    A blank cell is the usual way the sheet says "not bought yet", but a column
    pre-filled with formulas says it with a literal 0 instead. Importing those
    weeks would predate the position's first snapshot and, worse, hand the chart
    a zero first grid point -- which makes
    :func:`summa.portfolio.rebase_to_grid` return nothing and silently drop the
    benchmark line. A deposit without a value still starts the position: the
    money moved, so the week is real.
    """
    if value is not None and value > 0:
        return True
    return deposit is not None and deposit != 0


def build_snapshot_rows(
    columns: Sequence[PositionColumn],
    dates: Sequence[str],
    values: Sequence[Mapping[int, float | None]],
    deposits: Sequence[Mapping[int, float | None]],
    fx_rates: Mapping[str, float],
) -> tuple[list[SnapshotRow], int]:
    """Turn the parsed grid into snapshot rows.

    Once a position has started, every remaining week yields a row: a blank value
    is carried forward from the previous week and flagged ``carried``. The flag
    describes the value alone -- a deposit on such a week is real money and is
    recorded at its own date, because dropping it would distort ``invested_eur``
    and moving it would misdate the cash flow.

    A week before the position's first real value is not carried: there is nothing
    to carry. A deposit-only start therefore records ``carried=False``, the same as
    the literal-zero spelling of that week.

    :return: the rows plus the number of pre-start rows skipped.
    """
    rows: list[SnapshotRow] = []
    skipped: int = 0
    for position in columns:
        fx_rate: float = fx_rates.get(position.currency, DEFAULT_FX_RATE)
        started: bool = False
        seen_value: bool = False
        last_value: float = 0.0
        for index, snapshot_date in enumerate(dates):
            value: float | None = values[index].get(position.column)
            deposit: float | None = deposits[index].get(position.column)
            if not started:
                if not _position_started(value, deposit):
                    skipped += 1
                    continue
                started = True

            carried: bool = value is None and seen_value
            if value is not None:
                last_value = value
                seen_value = True
            rows.append(
                SnapshotRow(
                    depot=position.depot,
                    position=position.name,
                    date=snapshot_date,
                    value=last_value,
                    deposit=deposit if deposit is not None else 0.0,
                    fx_rate=fx_rate,
                    carried=carried,
                )
            )
    return rows, skipped


def parse_fx_overrides(pairs: Sequence[str]) -> dict[str, float]:
    """Parse repeated ``--fx CODE=RATE`` arguments.

    :raises ValueError: a pair is malformed or names an impossible rate.
    """
    overrides: dict[str, float] = {}
    for pair in pairs:
        raw_code, separator, raw_rate = pair.partition("=")
        if not separator:
            raise ValueError(f"--fx expects CODE=RATE, got {pair!r}")

        code: str = raw_code.strip().upper()
        if len(code) != CURRENCY_CODE_LENGTH or not (code.isascii() and code.isalpha()):
            raise ValueError(
                f"--fx needs a three-letter currency code, got {raw_code!r}"
            )

        try:
            rate: float = float(raw_rate.strip())
        except ValueError:
            raise ValueError(f"--fx rate must be a number, got {raw_rate!r}") from None
        if not math.isfinite(rate) or rate <= 0:
            raise ValueError(f"--fx rate must be greater than zero, got {raw_rate!r}")
        if code == DEFAULT_CURRENCY and rate != DEFAULT_FX_RATE:
            raise ValueError("--fx EUR is always 1.0 — EUR is the denominator")

        overrides[code] = rate
    return overrides


def format_summary(summary: ImportSummary) -> str:
    """Render the import summary as an indented block."""
    return "\n".join(
        [
            f"  depots     {summary.depots:5d}  ({summary.depots_created} created)",
            f"  positions  {summary.positions:5d}  ({summary.positions_created} created)",
            f"  snapshots  {summary.snapshots_written:5d} written, {summary.snapshots_existing} already present",
            f"  carried    {summary.rows_carried:5d}  rows copied forward",
            f"  skipped    {summary.rows_skipped:5d}  rows before a position's first value",
        ]
    )


# --- Workbook reading -------------------------------------------------------


def load_repaired_workbook(path: Path) -> Any:
    """Load a workbook, repacking it in memory so openpyxl accepts its styles.

    The file on disk is never touched. ``read_only`` is deliberately off: it
    hides ``merged_cells.ranges``, which is where the depot bands live.
    """
    buffer: io.BytesIO = io.BytesIO()
    with zipfile.ZipFile(path) as source:
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                data: bytes = source.read(item.filename)
                if item.filename == STYLES_PART:
                    data = strip_style_extensions(data)
                target.writestr(item, data)
    buffer.seek(0)
    return load_workbook(buffer, data_only=True)


def _read_text_row(worksheet: Any, row: int, last_column: int) -> dict[int, str]:
    """Return a row's non-empty text cells, keyed by column index."""
    texts: dict[int, str] = {}
    for column in range(1, last_column + 1):
        raw: Any = worksheet.cell(row=row, column=column).value
        if isinstance(raw, str) and raw.strip():
            texts[column] = raw.strip()
    return texts


def read_band_map(worksheet: Any, last_column: int) -> dict[int, str]:
    """Read row 1 and resolve every column to its depot band."""
    merges: list[tuple[int, int]] = []
    for merged in worksheet.merged_cells.ranges:
        if merged.min_row == BAND_ROW and merged.max_row == BAND_ROW:
            merges.append((int(merged.min_col), int(merged.max_col)))
    return fill_band_map(
        _read_text_row(worksheet, BAND_ROW, last_column), merges, last_column
    )


def read_grid(
    worksheet: Any, columns: Sequence[PositionColumn]
) -> tuple[list[str], list[dict[int, float | None]], list[dict[int, float | None]]]:
    """Read the data rows into dates plus per-column values and deposits.

    Rows whose date cell does not parse are skipped, which is how a trailing
    note or a blank spacer row below the data stays out of the import.
    """
    dates: list[str] = []
    values: list[dict[int, float | None]] = []
    deposits: list[dict[int, float | None]] = []
    for row in range(FIRST_DATA_ROW, int(worksheet.max_row) + 1):
        snapshot_date: str | None = parse_date_cell(
            worksheet.cell(row=row, column=DATE_COLUMN).value
        )
        if snapshot_date is None:
            continue

        row_values: dict[int, float | None] = {}
        row_deposits: dict[int, float | None] = {}
        for position in columns:
            row_values[position.column] = parse_number(
                worksheet.cell(row=row, column=position.column).value
            )
            row_deposits[position.column] = parse_number(
                worksheet.cell(row=row, column=position.column + 1).value
            )
        dates.append(snapshot_date)
        values.append(row_values)
        deposits.append(row_deposits)
    return dates, values, deposits


# --- Writing ----------------------------------------------------------------


def resolve_depot(
    cursor: sqlite3.Cursor, name: str, sort_order: int
) -> tuple[int, bool]:
    """Return a depot's id, creating it when missing.

    ``INSERT OR IGNORE ... RETURNING`` yields no row when it ignores, hence the
    separate lookup.

    :return: the id and whether this call created it.
    """
    cursor.execute(
        "INSERT OR IGNORE INTO portfolio_depots (name, sort_order) VALUES (?, ?)",
        (name, sort_order),
    )
    created: bool = cursor.rowcount == 1
    cursor.execute("SELECT id FROM portfolio_depots WHERE name = ?", (name,))
    row: Any = cursor.fetchone()
    return int(row["id"]), created


def resolve_position(
    cursor: sqlite3.Cursor, depot_id: int, position: PositionColumn
) -> tuple[int, bool]:
    """Return a position's id, creating it when missing.

    An existing position is never updated: the importer only guesses ``kind``
    and ``currency``, and the UI is allowed to have corrected that guess.
    """
    cursor.execute(
        "INSERT OR IGNORE INTO portfolio_positions "
        "(depot_id, name, kind, currency, sort_order) VALUES (?, ?, ?, ?, ?)",
        (
            depot_id,
            position.name,
            position.kind,
            position.currency,
            position.sort_order,
        ),
    )
    created: bool = cursor.rowcount == 1
    cursor.execute(
        "SELECT id FROM portfolio_positions WHERE depot_id = ? AND name = ?",
        (depot_id, position.name),
    )
    row: Any = cursor.fetchone()
    return int(row["id"]), created


def insert_snapshots(
    cursor: sqlite3.Cursor, position_id: int, rows: Sequence[SnapshotRow]
) -> tuple[int, int]:
    """Write a position's snapshots, leaving weeks already recorded untouched.

    Rows go in one at a time so ``rowcount`` can tell a write from a conflict;
    ``executemany`` would collapse the two.

    :return: how many rows were written and how many were already present.
    """
    written: int = 0
    for row in rows:
        cursor.execute(
            "INSERT OR IGNORE INTO portfolio_snapshots "
            "(position_id, date, value, deposit, fx_rate, carried) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (
                position_id,
                row.date,
                row.value,
                row.deposit,
                row.fx_rate,
                int(row.carried),
            ),
        )
        written += cursor.rowcount
    return written, len(rows) - written


def run_import(
    cursor: sqlite3.Cursor,
    columns: Sequence[PositionColumn],
    rows: Sequence[SnapshotRow],
    rows_skipped: int,
) -> ImportSummary:
    """Write depots, positions and snapshots, and report what happened."""
    rows_by_position: dict[tuple[str, str], list[SnapshotRow]] = {}
    for row in rows:
        rows_by_position.setdefault((row.depot, row.position), []).append(row)

    depot_ids: dict[str, int] = {}
    depots_created: int = 0
    positions_created: int = 0
    written: int = 0
    existing: int = 0
    carried: int = 0
    for position in columns:
        if position.depot not in depot_ids:
            depot_id, depot_is_new = resolve_depot(
                cursor, position.depot, len(depot_ids)
            )
            depot_ids[position.depot] = depot_id
            if depot_is_new:
                depots_created += 1

        position_id, position_is_new = resolve_position(
            cursor, depot_ids[position.depot], position
        )
        if position_is_new:
            positions_created += 1

        position_rows: list[SnapshotRow] = rows_by_position.get(
            (position.depot, position.name), []
        )
        row_written, row_existing = insert_snapshots(cursor, position_id, position_rows)
        written += row_written
        existing += row_existing
        carried += sum(1 for row in position_rows if row.carried)

    return ImportSummary(
        depots=len(depot_ids),
        depots_created=depots_created,
        positions=len(columns),
        positions_created=positions_created,
        snapshots_written=written,
        snapshots_existing=existing,
        rows_carried=carried,
        rows_skipped=rows_skipped,
    )


# --- Command line -----------------------------------------------------------


def _select_sheet(workbook: Any, sheet_name: str) -> Any:
    """Return the named sheet, falling back to the active one with a warning."""
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    fallback: Any = workbook.active
    print(
        f"warning: sheet {sheet_name!r} not found, reading {fallback.title!r} instead",
        file=sys.stderr,
    )
    return fallback


def _warn_unused_fx(
    columns: Sequence[PositionColumn], fx_rates: Mapping[str, float]
) -> None:
    """Warn about --fx rates no position's currency uses."""
    used: set[str] = {position.currency for position in columns}
    for code in sorted(set(fx_rates) - used):
        print(f"warning: --fx {code} matches no position currency", file=sys.stderr)


def import_workbook(
    path: Path,
    sheet_name: str,
    fx_rates: Mapping[str, float],
    database_path: Path,
    dry_run: bool,
) -> ImportSummary:
    """Read a workbook and write its history to the database.

    :raises ImportAbort: the sheet holds no recognisable position.
    """
    workbook: Any = load_repaired_workbook(path)
    worksheet: Any = _select_sheet(workbook, sheet_name)
    last_column: int = int(worksheet.max_column or 0)

    headers: dict[int, str] = _read_text_row(worksheet, HEADER_ROW, last_column)
    columns: list[PositionColumn] = find_position_columns(
        headers, read_band_map(worksheet, last_column)
    )
    if not columns:
        raise ImportAbort(
            f"no positions found in sheet {worksheet.title!r} — "
            f"row {HEADER_ROW} needs a {DEPOSIT_HEADER!r} column beside each position"
        )
    _warn_unused_fx(columns, fx_rates)

    dates, values, deposits = read_grid(worksheet, columns)
    rows, rows_skipped = build_snapshot_rows(columns, dates, values, deposits, fx_rates)

    conn: sqlite3.Connection = (
        connect_mirror(database_path) if dry_run else connect(database_path)
    )
    try:
        summary: ImportSummary = run_import(conn.cursor(), columns, rows, rows_skipped)
        # A dry run commits into its in-memory mirror, which close() discards.
        conn.commit()
        return summary
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def build_parser() -> argparse.ArgumentParser:
    """Build the command-line parser."""
    parser: argparse.ArgumentParser = argparse.ArgumentParser(
        prog="import_portfolio_xlsx",
        description="Import a depot workbook into Summa's portfolio tables.",
        epilog=(
            "Re-running is safe: a week already in the database is kept as it is, "
            "not refreshed, so a value corrected in the UI survives an import."
        ),
    )
    parser.add_argument("path", type=Path, help="the .xlsx workbook to read")
    parser.add_argument(
        "--db",
        type=Path,
        default=default_database_path(),
        help="database to write to (default: $DATABASE_PATH, else invoices.db)",
    )
    parser.add_argument(
        "--sheet",
        default=DEFAULT_SHEET,
        help=f"worksheet to read (default: {DEFAULT_SHEET})",
    )
    parser.add_argument(
        "--fx",
        action="append",
        default=[],
        metavar="CODE=RATE",
        help=(
            "exchange rate for a position currency, in units of that currency "
            "per EUR (e.g. --fx USD=1.08). Repeatable; currencies without a "
            "rate are stored at 1.0."
        ),
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="report what would be written without writing it",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    """Run the importer and return a process exit code."""
    args: argparse.Namespace = build_parser().parse_args(argv)
    path: Path = args.path

    try:
        fx_rates: dict[str, float] = parse_fx_overrides(args.fx)
    except ValueError as error:
        print(f"error: {error}", file=sys.stderr)
        return EXIT_USAGE_ERROR
    if not path.is_file():
        print(f"error: no such workbook: {path}", file=sys.stderr)
        return EXIT_USAGE_ERROR

    try:
        summary: ImportSummary = import_workbook(
            path, args.sheet, fx_rates, args.db, dry_run=args.dry_run
        )
    except (ImportAbort, OSError, zipfile.BadZipFile, sqlite3.Error) as error:
        print(f"error: {error}", file=sys.stderr)
        return EXIT_IMPORT_ERROR

    headline: str = "Dry run — nothing written" if args.dry_run else "Imported"
    print(f"{headline}: {path} -> {args.db}")
    print(format_summary(summary))
    return EXIT_OK


if __name__ == "__main__":
    raise SystemExit(main())
